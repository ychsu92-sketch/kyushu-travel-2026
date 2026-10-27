#!/usr/bin/env python3
"""
生活小秘書 - Google 試算表自動同步腳本
用法: python3 sync_from_sheets.py
自動從 Google Sheets 下載最新南九州行程與費用，並重新編譯儀表板。
"""

import os
import json
import urllib.request
import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, 'data')
SHEET_ID = '1P4niaKfXpHnrlVkiexFmxa2Q0U7kwOvj6yUINn--WNo'
XLSX_PATH = os.path.join(BASE_DIR, 'travel_data_import.xlsx')

def download_sheet():
    url = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
    headers = {'User-Agent': 'Mozilla/5.0'}
    req = urllib.request.Request(url, headers=headers)
    try:
        with urllib.request.urlopen(req, timeout=10) as resp:
            data = resp.read()
            with open(XLSX_PATH, 'wb') as f:
                f.write(data)
            print(f"📥 成功下載最新 Google 試算表 ({len(data)} bytes)")
            return True
    except Exception as e:
        print(f"⚠️ 無法下載線上試算表 (將使用現有本地資料): {e}")
        return False

def sync():
    download_sheet()
    if not os.path.exists(XLSX_PATH):
        print("❌ 找不到 travel_data_import.xlsx")
        return

    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    
    # 讀取 行程表
    ws_itinerary = wb['行程表']
    days = []
    for col_idx in range(2, 9):
        day_col_header = ws_itinerary.cell(5, col_idx).value
        day_data = {'col_header': day_col_header}
        for row_idx in range(6, 23):
            row_title = ws_itinerary.cell(row_idx, 1).value
            cell_val = ws_itinerary.cell(row_idx, col_idx).value
            if row_title:
                day_data[row_title.strip()] = cell_val
        days.append(day_data)

    # 讀取 費用
    ws_cost = wb['費用']
    expenses = []
    for r in range(18, 42):
        row_vals = [ws_cost.cell(r, c).value for c in range(2, 16)]
        if any(row_vals[:6]) and row_vals[1]:
            expenses.append({
                'date': str(row_vals[0] or ''),
                'item': str(row_vals[1] or ''),
                'category': str(row_vals[2] or ''),
                'payer': str(row_vals[3] or ''),
                'currency': str(row_vals[4] or ''),
                'amount_original': row_vals[5],
                'amount_twd': row_vals[6],
                'achun': bool(row_vals[7]),
                'alu': bool(row_vals[8]),
                'aunt': bool(row_vals[9]),
                'uncle': bool(row_vals[10]),
                'headcount': row_vals[11],
                'per_person_twd': row_vals[12],
                'note': str(row_vals[13] or '')
            })

    # 儲存中繼檔
    with open(os.path.join(DATA_DIR, 'parsed_kyushu_trip.json'), 'w', encoding='utf-8') as f:
        json.dump({'days': days, 'expenses': expenses}, f, ensure_ascii=False, indent=2)

    # 同步更新 current_trip.json
    current_trip_path = os.path.join(DATA_DIR, 'current_trip.json')
    if os.path.exists(current_trip_path):
        with open(current_trip_path, 'r', encoding='utf-8') as f:
            trip = json.load(f)

        items = []
        total_twd = 0
        for e in expenses:
            twd_val = round(float(e['amount_twd'] or 0), 2)
            total_twd += twd_val
            items.append({
                'day': e['date'],
                'item': e['item'],
                'category': e['category'],
                'payer': e['payer'],
                'currency': e['currency'],
                'amount_orig': float(e['amount_original'] or twd_val),
                'amount_twd': twd_val,
                'split': {
                    'achun': e['achun'],
                    'alu': e['alu'],
                    'ayee': e['aunt'],
                    'yijhang': e['uncle']
                }
            })
            if '華航' in e['item'] or '機票' in e['item']:
                if 'flights' in trip and 'outbound' in trip['flights']:
                    trip['flights']['outbound']['flight_cost_twd'] = int(twd_val)

        trip['expenses_list'] = items
        trip['expense_summary'] = {
            'total_spent_twd': total_twd,
            'per_person_avg_twd': round(total_twd / 4, 1),
            'status': '已平帳 (收支平衡)'
        }

        with open(current_trip_path, 'w', encoding='utf-8') as f:
            json.dump(trip, f, ensure_ascii=False, indent=2)
        print(f"✅ current_trip.json 已成功同步試算表共 {len(items)} 筆公費支出，總額 NT$ {total_twd:,.0f}")

    # 執行建置
    import build_dashboard
    build_dashboard.build()
    print("✨ 同步完成，儀表板已更新！")

if __name__ == '__main__':
    sync()
